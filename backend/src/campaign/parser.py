"""Spreadsheet and document parsing for campaigns.

Parses recipients files (CSV/XLSX), schedule files (CSV/XLSX),
and email document files (Markdown, DOCX).
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time

import markdown
import openpyxl
from fastapi import UploadFile

EMAIL_RE = re.compile(r"^[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}$")
REQUIRED_RECIPIENT_COLS = {"recipient_email", "recipient_name"}
REQUIRED_SCHEDULE_COLS = {"send_date", "send_time", "email_template_ref"}


@dataclass
class RecipientData:
    email: str
    name: str | None
    custom_fields: dict


@dataclass
class RecipientsParseResult:
    recipients: list[RecipientData] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class ScheduleStepData:
    step_order: int
    send_date: date | None
    relative_days: int | None
    send_time: time
    email_template_ref: str


@dataclass
class ScheduleParseResult:
    steps: list[ScheduleStepData] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class EmailTemplate:
    ref: str
    subject: str
    body_html: str
    body_text: str


# ── Helpers ───────────────────────────────────────────────────────────────


def _read_rows_from_upload(
    file: UploadFile,
) -> list[list[str]]:
    """Read rows from a CSV or XLSX upload, returning list of string rows."""
    content = file.file.read()
    filename = file.filename or ""

    if filename.endswith(".xlsx"):
        return _read_xlsx_rows(content)
    elif filename.endswith(".csv"):
        return _read_csv_rows(content)
    else:
        # Try to detect from content
        try:
            return _read_xlsx_rows(content)
        except Exception:
            return _read_csv_rows(content)


def _read_xlsx_rows(content: bytes) -> list[list[str]]:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(cell).strip() if cell is not None else "" for cell in row])
    wb.close()
    return rows


def _read_csv_rows(content: bytes) -> list[list[str]]:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    return [[cell.strip() for cell in row] for row in reader]


def _is_empty_row(row: list[str]) -> bool:
    return all(cell == "" for cell in row)


def _parse_date(value: str) -> tuple[date | None, int | None, str | None]:
    """Parse a send_date value.

    Returns (absolute_date, relative_days, error).
    """
    value = value.strip()
    if not value:
        return None, None, "send_date is empty"

    # Check if it's a relative integer
    try:
        days = int(value)
        if days < 0:
            return None, None, f"Relative days must be non-negative, got {days}"
        return None, days, None
    except ValueError:
        pass

    # Try absolute date
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date(), None, None
        except ValueError:
            continue

    return None, None, f"Invalid date format: '{value}'"


def _parse_time(value: str) -> tuple[time | None, str | None]:
    """Parse a send_time value. Returns (time, error)."""
    value = value.strip()
    if not value:
        return None, "send_time is empty"

    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p"):
        try:
            return datetime.strptime(value, fmt).time(), None
        except ValueError:
            continue

    return None, f"Invalid time format: '{value}'"


# ── Recipients Parser ─────────────────────────────────────────────────────


def parse_recipients_file(file: UploadFile) -> RecipientsParseResult:
    """Parse a recipients CSV or XLSX file."""
    result = RecipientsParseResult()

    try:
        rows = _read_rows_from_upload(file)
    except Exception as e:
        result.errors.append(f"Failed to read file: {e}")
        return result

    if not rows:
        result.errors.append("File is empty.")
        return result

    headers = [h.lower().strip() for h in rows[0]]

    missing = REQUIRED_RECIPIENT_COLS - set(headers)
    if missing:
        result.errors.append(f"Missing required columns: {', '.join(sorted(missing))}")
        return result

    email_idx = headers.index("recipient_email")
    name_idx = headers.index("recipient_name")
    custom_indices = {
        headers[i]: i for i in range(len(headers)) if headers[i] not in REQUIRED_RECIPIENT_COLS
    }

    for row_num, row in enumerate(rows[1:], start=2):
        if _is_empty_row(row):
            continue

        # Pad row if shorter than headers
        while len(row) < len(headers):
            row.append("")

        email_val = row[email_idx].strip()
        name_val = row[name_idx].strip() or None

        if not email_val:
            result.errors.append(f"Row {row_num}: recipient_email is empty")
            continue

        if not EMAIL_RE.match(email_val):
            result.errors.append(f"Row {row_num}: invalid email address '{email_val}'")
            continue

        custom_fields = {}
        for col_name, idx in custom_indices.items():
            val = row[idx].strip()
            if val:
                custom_fields[col_name] = val

        result.recipients.append(
            RecipientData(
                email=email_val,
                name=name_val,
                custom_fields=custom_fields if custom_fields else {},
            )
        )

    return result


# ── Schedule Parser ───────────────────────────────────────────────────────


def parse_schedule_file(file: UploadFile) -> ScheduleParseResult:
    """Parse a schedule CSV or XLSX file."""
    result = ScheduleParseResult()

    try:
        rows = _read_rows_from_upload(file)
    except Exception as e:
        result.errors.append(f"Failed to read file: {e}")
        return result

    if not rows:
        result.errors.append("File is empty.")
        return result

    headers = [h.lower().strip() for h in rows[0]]

    missing = REQUIRED_SCHEDULE_COLS - set(headers)
    if missing:
        result.errors.append(f"Missing required columns: {', '.join(sorted(missing))}")
        return result

    date_idx = headers.index("send_date")
    time_idx = headers.index("send_time")
    ref_idx = headers.index("email_template_ref")

    today = date.today()

    for row_num, row in enumerate(rows[1:], start=2):
        if _is_empty_row(row):
            continue

        while len(row) < len(headers):
            row.append("")

        step_order = len(result.steps) + 1

        # Parse date
        abs_date, rel_days, date_err = _parse_date(row[date_idx])
        if date_err:
            result.errors.append(f"Row {row_num}: {date_err}")
            continue

        # First step must be absolute
        if step_order == 1 and rel_days is not None:
            result.errors.append(
                f"Row {row_num}: first schedule step must have "
                f"an absolute date, not a relative value"
            )
            continue

        # Warn on past absolute dates
        if abs_date and abs_date < today:
            result.warnings.append(f"Row {row_num}: send_date {abs_date} is in the past")

        # Parse time
        send_time, time_err = _parse_time(row[time_idx])
        if time_err:
            result.errors.append(f"Row {row_num}: {time_err}")
            continue

        ref = row[ref_idx].strip()
        if not ref:
            result.errors.append(f"Row {row_num}: email_template_ref is empty")
            continue

        result.steps.append(
            ScheduleStepData(
                step_order=step_order,
                send_date=abs_date,
                relative_days=rel_days,
                send_time=send_time,
                email_template_ref=ref,
            )
        )

    return result


# ── Email Document Parser ─────────────────────────────────────────────────


def parse_email_document(file: UploadFile) -> list[EmailTemplate]:
    """Parse an email document file (.md or .docx) into templates.

    Format: sections separated by '---', each with frontmatter lines
    (ref: ..., subject: ...) followed by body content.
    """
    content = file.file.read()
    filename = file.filename or ""

    if filename.endswith(".docx"):
        text = _read_docx_text(content)
    else:
        text = content.decode("utf-8-sig")

    return _parse_email_text(text)


def _read_docx_text(content: bytes) -> str:
    """Extract text from a DOCX file."""
    import docx

    doc = docx.Document(io.BytesIO(content))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_email_text(text: str) -> list[EmailTemplate]:
    """Parse email text content into templates."""
    sections = re.split(r"\n---\s*\n", text.strip())
    templates: list[EmailTemplate] = []

    for i, section in enumerate(sections):
        section = section.strip()
        if not section:
            continue

        ref = str(i + 1)
        subject = ""
        body_lines: list[str] = []
        in_frontmatter = True

        for line in section.split("\n"):
            if in_frontmatter:
                ref_match = re.match(r"^ref:\s*(.+)$", line, re.IGNORECASE)
                subject_match = re.match(r"^subject:\s*(.+)$", line, re.IGNORECASE)
                if ref_match:
                    ref = ref_match.group(1).strip()
                    continue
                elif subject_match:
                    subject = subject_match.group(1).strip()
                    continue
                else:
                    in_frontmatter = False

            body_lines.append(line)

        body_text = "\n".join(body_lines).strip()
        body_html = markdown.markdown(body_text)

        if subject or body_text:
            templates.append(
                EmailTemplate(
                    ref=ref,
                    subject=subject,
                    body_html=body_html,
                    body_text=body_text,
                )
            )

    return templates
